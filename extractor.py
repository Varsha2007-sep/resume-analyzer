"""Resume text extraction for PDF and DOCX files."""
import io
import re

from docx import Document
from pypdf import PdfReader

ALLOWED = {"pdf", "docx"}
MAX_BYTES = 5 * 1024 * 1024


class ExtractionError(Exception):
    pass


def _clean(text: str) -> str:
    text = text.replace("\r", "\n").replace("\u00a0", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _pdf(data: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ExtractionError("This PDF is password protected. Please upload an unlocked copy.")
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    except ExtractionError:
        raise
    except Exception as exc:  # corrupted / unsupported
        raise ExtractionError("Could not read this PDF. It may be corrupted.") from exc


def _docx(data: bytes) -> str:
    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError("Could not read this DOCX file. It may be corrupted.") from exc
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:  # many templates put skills/education in tables
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def extract_text(filename: str, data: bytes) -> str:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in ALLOWED:
        raise ExtractionError("Unsupported file type. Please upload a PDF or DOCX resume.")
    if not data:
        raise ExtractionError("The uploaded file is empty.")
    if len(data) > MAX_BYTES:
        raise ExtractionError("File is larger than 5 MB.")
    text = _clean(_pdf(data) if ext == "pdf" else _docx(data))
    if len(text) < 80:
        raise ExtractionError("Very little text could be extracted. If this is a scanned/image PDF, "
                              "export a text-based PDF or upload a DOCX instead.")
    return text
